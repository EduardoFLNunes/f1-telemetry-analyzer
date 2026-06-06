# -*- coding: utf-8 -*-
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Manual_Tecnico_F1_Telemetry_Analyzer_ABNT.docx"


BLUE = RGBColor(31, 78, 121)
MUTED = RGBColor(89, 89, 89)
BLACK = RGBColor(0, 0, 0)
LIGHT_FILL = "E8EEF5"
SOFT_FILL = "F4F6F9"


def set_run(run, *, size=12, bold=False, italic=False, color=BLACK, font="Arial"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_paragraph_spacing(paragraph, *, before=0, after=6, line=1.5, first_line=True):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if first_line:
        fmt.first_line_indent = Cm(1.25)
    else:
        fmt.first_line_indent = Cm(0)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, *, bold=False, color=BLACK, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.alignment = align
    set_paragraph_spacing(p, after=0, line=1.15, first_line=False)
    run = p.add_run(str(text))
    set_run(run, size=size, bold=bold, color=color)


def add_page_field(paragraph, field_code="PAGE"):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Página ")
    set_run(run, size=9, color=MUTED)
    add_page_field(footer)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(1.25)

    for name, size, color, before, after in (
        ("Heading 1", 12, BLUE, 16, 8),
        ("Heading 2", 12, BLUE, 12, 6),
        ("Heading 3", 12, RGBColor(47, 84, 150), 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.first_line_indent = Cm(0)

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(11)
        style.paragraph_format.line_spacing = 1.3
        style.paragraph_format.space_after = Pt(4)


def add_cover(doc: Document):
    for text, size, bold, after in (
        ("F1 TELEMETRY ANALYZER", 14, True, 8),
        ("MANUAL TÉCNICO DO SISTEMA", 14, True, 2),
        ("Documentação técnica conforme estrutura e apresentação ABNT", 11, False, 0),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, before=0, after=after, line=1.15, first_line=False)
        run = p.add_run(text)
        set_run(run, size=size, bold=bold, color=BLACK)

    for _ in range(10):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=6, line=1.15, first_line=False)
    r = p.add_run("Sistema de análise de telemetria, Racing Line, comparação de voltas,\nmodelo físico, coach de corrida e integração com Assetto Corsa")
    set_run(r, size=12, bold=True, color=BLUE)

    for _ in range(12):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=2, line=1.15, first_line=False)
    set_run(p.add_run("São Paulo"), size=12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=0, line=1.15, first_line=False)
    set_run(p.add_run("2026"), size=12)
    doc.add_page_break()


def add_title_page(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=18, line=1.15, first_line=False)
    set_run(p.add_run("F1 TELEMETRY ANALYZER"), size=12, bold=True)

    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=10, line=1.15, first_line=False)
    set_run(p.add_run("MANUAL TÉCNICO DO SISTEMA"), size=14, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(7)
    set_paragraph_spacing(p, after=8, line=1.5, first_line=False)
    text = (
        "Manual técnico elaborado para explicar, de forma clara e completa, "
        "a arquitetura, os fluxos de dados, as integrações, os módulos de análise "
        "e os procedimentos de uso do sistema F1 Telemetry Analyzer no estado atual "
        "do projeto."
    )
    set_run(p.add_run(text), size=12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(7)
    set_paragraph_spacing(p, after=8, line=1.5, first_line=False)
    set_run(p.add_run("Elaboração: documentação gerada a partir do repositório local do sistema."), size=12)

    for _ in range(8):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=2, line=1.15, first_line=False)
    set_run(p.add_run("São Paulo"), size=12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, after=0, line=1.15, first_line=False)
    set_run(p.add_run("2026"), size=12)
    doc.add_page_break()


def heading(doc, level: int, text: str):
    p = doc.add_paragraph(style=f"Heading {level}")
    label = text.upper() if level == 1 else text
    run = p.add_run(label)
    set_run(run, size=12, bold=True, color=BLUE if level <= 2 else RGBColor(47, 84, 150))
    return p


def para(doc, text: str, *, no_indent=False, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, after=6, line=1.5, first_line=not no_indent)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run(r, bold=True)
        rest = text[len(bold_prefix):]
        if rest:
            set_run(p.add_run(rest))
    else:
        set_run(p.add_run(text))
    return p


def bullet(doc, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, after=4, line=1.3, first_line=False)
    set_run(p.add_run(text), size=11)


def numbered(doc, text: str):
    p = doc.add_paragraph(style="List Number")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(p, after=4, line=1.3, first_line=False)
    set_run(p.add_run(text), size=11)


def caption(doc, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=6, after=3, line=1.15, first_line=False)
    set_run(p.add_run(text), size=10, bold=True)


def source_note(doc, text="Fonte: elaborado a partir do código-fonte do projeto."):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, before=3, after=6, line=1.15, first_line=False)
    set_run(p.add_run(text), size=9, italic=True, color=MUTED)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        shade_cell(hdr[i], LIGHT_FILL)
        set_cell_text(hdr[i], h, bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        if widths:
            hdr[i].width = Cm(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=9.2)
            if widths:
                cells[i].width = Cm(widths[i])
    return table


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade_cell(cell, SOFT_FILL)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=0, line=1.05, first_line=False)
    run = p.add_run("\n".join(lines))
    set_run(run, size=9, font="Courier New", color=RGBColor(32, 32, 32))
    return table


def add_summary(doc: Document):
    heading(doc, 1, "Resumo")
    para(
        doc,
        "Este manual descreve o sistema F1 Telemetry Analyzer no estado atual do desenvolvimento. "
        "A aplicação combina backend FastAPI, frontend React, telemetria em tempo real, integração "
        "por memória compartilhada com Assetto Corsa, geometria de pista obtida de arquivos KN5/AI, "
        "Racing Line de referência baseada na melhor volta válida, comparação por microsetores, "
        "modelo físico simplificado de perfil de velocidade e painéis de apoio à pilotagem.",
    )
    para(
        doc,
        "O foco deste documento é permitir que uma pessoa que não participou do desenvolvimento "
        "entenda o que o sistema faz, como os dados circulam, como executar a aplicação, quais "
        "módulos compõem a solução, quais riscos já foram tratados e quais limitações ainda existem.",
    )
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=8, line=1.5, first_line=False)
    set_run(p.add_run("Palavras-chave: "), bold=True)
    set_run(p.add_run("telemetria; automobilismo; Assetto Corsa; Racing Line; FastAPI; React."))

    heading(doc, 1, "Lista de siglas")
    caption(doc, "Quadro 1 - Siglas e significados")
    add_table(
        doc,
        ["Sigla", "Significado no sistema"],
        [
            ["ABNT", "Associação Brasileira de Normas Técnicas."],
            ["AC", "Assetto Corsa, simulador de corrida usado como fonte real de telemetria."],
            ["API", "Interface de programação usada pelo frontend para consultar o backend."],
            ["CSV", "Arquivo tabular usado para telemetria gravada ou dados de pista legados."],
            ["KN5", "Formato de malha 3D de pistas/carros do Assetto Corsa."],
            ["RL", "Racing Line, linha de referência usada para comparação do piloto."],
            ["UI", "Interface de usuário no navegador."],
            ["UDP", "Protocolo usado para receber dados de oponentes e capturas de outros jogos."],
            ["WS", "WebSocket usado para transmitir telemetria ao frontend."],
        ],
        widths=[2.8, 12.0],
    )
    source_note(doc)
    doc.add_page_break()


def add_static_toc(doc: Document):
    heading(doc, 1, "Sumário")
    items = [
        "1 Introdução",
        "2 Visão geral do sistema",
        "3 Requisitos e tecnologias",
        "4 Estrutura do projeto",
        "5 Arquitetura do backend",
        "6 Pipeline de telemetria",
        "7 Racing Line e análise de microsetores",
        "8 Physics Fast Lap",
        "9 Inteligência de pilotagem e painéis de apoio",
        "10 Arquitetura do frontend",
        "11 Oponentes, gravação e persistência",
        "12 Execução operacional",
        "13 API do sistema",
        "14 Testes, validação e baseline",
        "15 Troubleshooting",
        "16 Limitações, decisões técnicas e roadmap",
        "Referências",
        "Apêndice A - Variáveis de ambiente",
        "Apêndice B - Glossário técnico",
    ]
    for item in items:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, after=2, line=1.2, first_line=False)
        set_run(p.add_run(item), size=11)
    source_note(doc, "Fonte: organização do próprio manual.")
    doc.add_page_break()


def add_body(doc: Document):
    heading(doc, 1, "1 Introdução")
    heading(doc, 2, "1.1 Finalidade")
    para(
        doc,
        "O F1 Telemetry Analyzer é uma estação de análise de pilotagem para simulação de corrida. "
        "Ele lê telemetria do carro, projeta a posição na geometria da pista, compara voltas, "
        "identifica perdas por trecho e apresenta diagnósticos visuais no mapa. A meta prática é "
        "dar ao piloto uma leitura rápida de onde freia, acelera, posiciona o carro e perde tempo.",
    )
    para(
        doc,
        "A primeira Racing Line do sistema não é uma linha perfeita matemática. Ela é uma linha "
        "de referência real, retirada da melhor volta válida disponível. Essa decisão torna a "
        "ferramenta útil mais cedo: o piloto compara seu comportamento atual com uma volta real "
        "já realizada antes de evoluir para uma linha ideal composta.",
    )
    heading(doc, 2, "1.2 Público-alvo")
    bullet(doc, "Pilotos de simulador que querem entender sua volta com leitura visual e numérica.")
    bullet(doc, "Desenvolvedores que precisam manter ou expandir o backend, o frontend e as integrações.")
    bullet(doc, "Usuários técnicos que precisam operar o sistema com Assetto Corsa e dados gravados.")
    heading(doc, 2, "1.3 Escopo")
    para(
        doc,
        "Este manual cobre o estado atual do sistema até 2 de junho de 2026. Ele descreve módulos "
        "existentes, endpoints, fluxos de dados, formatos, modos de execução, testes e limitações. "
        "Não cobre instalação avançada de mods do Assetto Corsa nem substitui a documentação oficial "
        "das normas ABNT; apenas aplica as convenções de apresentação, estrutura e referência cabíveis.",
    )

    heading(doc, 1, "2 Visão geral do sistema")
    heading(doc, 2, "2.1 O que o sistema faz")
    bullet(doc, "Recebe telemetria ao vivo do Assetto Corsa ou usa fontes mock/replay para desenvolvimento.")
    bullet(doc, "Carrega ou reconstrói a geometria da pista e mantém a pista em um espaço canônico X/Z.")
    bullet(doc, "Projeta o carro na pista e calcula distância percorrida, offset lateral e progresso da volta.")
    bullet(doc, "Seleciona a melhor volta válida como Racing Line de referência.")
    bullet(doc, "Compara volta atual, referência e oponentes em microsetores configuráveis.")
    bullet(doc, "Exibe a Racing Line como linha contínua e, opcionalmente, como gradiente de velocidade/freio.")
    bullet(doc, "Gera relatório de Physics Fast Lap sobre a Racing Line existente.")
    bullet(doc, "Mostra painéis de física, comparação, coach, telemetria, oponentes e performance.")

    heading(doc, 2, "2.2 Arquitetura em camadas")
    caption(doc, "Figura 1 - Fluxo lógico do sistema")
    add_code_block(
        doc,
        [
            "Assetto Corsa / Replay / Mock",
            "        |",
            "        v",
            "TelemetrySourceManager -> TelemetryRuntime -> RuntimeState",
            "        |                         |",
            "        |                         v",
            "        |                  ProjectionEngine + TrackGeometry",
            "        |                         |",
            "        v                         v",
            "WebSocket / REST API -> React + Zustand -> Canvas TrackRenderer",
            "        |                         |",
            "        v                         v",
            "Racing Line, Comparison, Coach, Physics Fast Lap, Recording",
        ],
    )
    source_note(doc)

    heading(doc, 2, "2.3 Princípios de projeto")
    bullet(doc, "A geometria da pista é tratada como referência estável. O carro e os overlays usam o mesmo espaço de mapa.")
    bullet(doc, "O overlay da Racing Line não tem projeção própria; ele reutiliza a transformação mundo X/Z para mapa com Y = -Z.")
    bullet(doc, "A linha visual principal usa amostras reais da volta de referência, sem suavização agressiva que corte curvas.")
    bullet(doc, "Microsetores são usados para análise e tabela, não para desenhar a linha principal como pista segmentada.")
    bullet(doc, "O backend inicia em modo seguro sem auto-conectar ao Assetto, evitando criar memória compartilhada antes do simulador.")
    bullet(doc, "Os cálculos pesados e payloads grandes são desacoplados por cache, polling controlado e modos de performance.")

    heading(doc, 1, "3 Requisitos e tecnologias")
    heading(doc, 2, "3.1 Requisitos de execução")
    caption(doc, "Quadro 2 - Requisitos principais")
    add_table(
        doc,
        ["Item", "Descrição"],
        [
            ["Sistema operacional", "Windows é o ambiente principal validado; Linux/macOS são suportados para fluxos sem Assetto Corsa."],
            ["Python", "Python 3.9+; no workspace atual há ambiente virtual em .venv."],
            ["Node.js", "Node.js 18+ e npm para o frontend React/Vite."],
            ["Assetto Corsa", "Necessário para telemetria real por memória compartilhada."],
            ["Navegador", "Frontend em http://localhost:5173."],
            ["Backend", "FastAPI em http://localhost:8000."],
        ],
        widths=[4.1, 10.7],
    )
    source_note(doc)

    heading(doc, 2, "3.2 Tecnologias do backend")
    caption(doc, "Quadro 3 - Bibliotecas Python")
    add_table(
        doc,
        ["Biblioteca", "Uso"],
        [
            ["FastAPI e Uvicorn", "API REST, WebSocket e servidor ASGI."],
            ["Pandas, NumPy e SciPy", "Leitura de CSV, cálculo vetorial, estatística e interpolação."],
            ["Pydantic", "Estruturas e validação de dados."],
            ["DuckDB e PyArrow", "Persistência local e dados tabulares/colunares."],
            ["FastF1", "Integração histórica com dados reais de F1."],
            ["ONNX Runtime, scikit-learn, PyTorch, Transformers", "Base para módulos de IA, modelos e experimentos futuros."],
            ["Gymnasium e NetworkX", "Ambiente de RL e grafos de conhecimento/memória."],
        ],
        widths=[4.8, 10.0],
    )
    source_note(doc, "Fonte: backend/requirements.txt.")

    heading(doc, 2, "3.3 Tecnologias do frontend")
    caption(doc, "Quadro 4 - Bibliotecas JavaScript/TypeScript")
    add_table(
        doc,
        ["Biblioteca", "Uso"],
        [
            ["React 18", "Composição da interface e painéis."],
            ["Vite", "Servidor de desenvolvimento e build."],
            ["Zustand", "Store global de telemetria e estado da UI."],
            ["Axios", "Cliente HTTP para endpoints REST."],
            ["Lucide React", "Ícones dos controles e painéis."],
            ["Plotly/Recharts", "Gráficos e componentes históricos ainda disponíveis."],
            ["TypeScript", "Tipos do store, painéis e payloads principais."],
        ],
        widths=[4.8, 10.0],
    )
    source_note(doc, "Fonte: frontend/package.json.")

    heading(doc, 1, "4 Estrutura do projeto")
    caption(doc, "Quadro 5 - Diretórios e responsabilidades")
    add_table(
        doc,
        ["Caminho", "Responsabilidade"],
        [
            ["backend/main.py", "Aplicação FastAPI, endpoints, ciclo de vida, runtime e orquestração geral."],
            ["backend/core", "Módulos de telemetria, pista, física, análise, IA, gravação e integração com Assetto."],
            ["backend/tests", "Testes unitários dos módulos críticos."],
            ["frontend/src", "Aplicação React, store, componentes e cliente de API."],
            ["frontend/src/components/map", "Renderização do mapa, carro, oponentes, pista e Racing Line em canvas."],
            ["data/cache/tracks", "Geometrias de pista cacheadas, especialmente Interlagos KN5."],
            ["data/telemetry_db e data/feature_store", "Bases DuckDB locais para telemetria e features."],
            ["tools/assetto_opponents_exporter", "Aplicativo Python do Assetto para exportar oponentes via UDP."],
            ["outputs", "Artefatos de validação, logs e capturas geradas durante desenvolvimento."],
        ],
        widths=[5.4, 9.4],
    )
    source_note(doc)

    heading(doc, 1, "5 Arquitetura do backend")
    heading(doc, 2, "5.1 Aplicação FastAPI")
    para(
        doc,
        "O arquivo backend/main.py centraliza a aplicação. Ele cria objetos globais de runtime, cache, "
        "buffer de telemetria, estado da pista, runtime de oponentes, gravação e runtime cognitivo. "
        "Durante o lifespan, o sistema seleciona a fonte de telemetria, inicializa a geometria da pista, "
        "inicia a thread de leitura e registra eventos para WebSocket, gravação e painéis.",
    )
    heading(doc, 2, "5.2 Fontes de telemetria")
    para(
        doc,
        "O TelemetrySourceManager escolhe entre mock, replay CSV e Assetto Corsa. O modo auto agora "
        "não tenta conectar ao Assetto por padrão. Essa escolha foi feita para impedir que a aplicação "
        "crie ou segure páginas de memória compartilhada antes do simulador estar pronto.",
    )
    caption(doc, "Quadro 6 - Fontes de telemetria")
    add_table(
        doc,
        ["Fonte", "Quando usar", "Comportamento"],
        [
            ["mock", "Desenvolvimento e inicialização segura.", "Não lê AC e mantém a aplicação estável com geometria padrão."],
            ["replay", "Reproduzir CSV de telemetria.", "Carrega amostras de data/example_telemetry.csv ou outro fixture."],
            ["assetto_corsa", "Sessão real aberta na pista.", "Exige leitura válida antes de trocar a fonte; se não houver amostra, falha com mensagem clara."],
            ["auto", "Inicialização geral.", "Com ASSETTO_AUTO_CONNECT=0 seleciona mock; com 1 pode promover para Assetto quando houver amostra real."],
        ],
        widths=[3.0, 5.2, 6.6],
    )
    source_note(doc, "Fonte: backend/core/telemetry/telemetry_reader_impl.py.")

    heading(doc, 2, "5.3 Integração segura com Assetto Corsa")
    para(
        doc,
        "A integração com Assetto usa as páginas de memória compartilhada acpmf_physics, acpmf_graphics "
        "e acpmf_static. O Assetto deve ser o dono dessas páginas. Por isso, o adaptador primeiro chama "
        "OpenFileMappingW para verificar se elas já existem; só depois abre mmap em modo leitura. "
        "Se alguma página não existir, a conexão retorna falso e não cria nomes de memória por acidente.",
    )
    para(
        doc,
        "Fluxo operacional recomendado: abrir o Assetto Corsa, entrar em uma sessão de pilotagem, iniciar "
        "ou manter o backend em modo seguro, conectar manualmente a fonte assetto_corsa e, antes de fechar "
        "ou reiniciar o Assetto, voltar a fonte para mock para liberar handles.",
    )

    heading(doc, 2, "5.4 Estado em tempo real e projeção")
    para(
        doc,
        "O RuntimeState guarda a pista ativa, o último sample, o estado projetado do carro e o motor de "
        "projeção. Cada TelemetrySample é convertido para uma estrutura de UI contendo worldPosition, "
        "mapPosition, projectedPosition, distanceAlongTrack, lateralOffset, velocidade, pedais, marcha, "
        "RPM e dados de física. Essa estrutura alimenta WebSocket e endpoints REST.",
    )
    para(
        doc,
        "A convenção espacial é fundamental: a amostra canônica usa coordenadas de mundo X/Z. Para o mapa, "
        "o eixo visual Y é calculado como -Z. O carro, a pista, os oponentes e a Racing Line precisam "
        "usar essa mesma convenção para evitar espelhamento ou desalinhamento.",
    )

    heading(doc, 2, "5.5 Geometria de pista")
    para(
        doc,
        "A geometria principal de Interlagos vem dos arquivos do Assetto Corsa. O provider KN5 Surface "
        "Track Geometry lê a malha da pista, cruza com fast_lane.ai e calcula bordas por intervalos de "
        "superfície. O cache atual de vhe_interlagos gp possui aproximadamente 4.345,88 m, loop fechado, "
        "2.680 amostras válidas de raycast e largura média aproximada de 12,50 m.",
    )
    para(
        doc,
        "Quando a fonte é mock, o backend carrega a geometria fixa de Interlagos para permitir uso e testes "
        "sem Assetto. Quando a fonte é Assetto, o sistema tenta carregar KN5/cache; reconstrução por "
        "trajetória do piloto fica desabilitada para evitar que uma volta imperfeita vire geometria oficial.",
    )

    heading(doc, 2, "5.6 Coleta de volta e buffer")
    para(
        doc,
        "O TelemetryBuffer mantém até 20.000 amostras. O LapCollector detecta o fechamento da volta pelo "
        "wrap de normalizedSplinePosition, de valores altos próximos a 0,95 para baixos próximos a 0,10. "
        "Uma volta candidata precisa ter amostras suficientes, distância plausível e área de pista mínima.",
    )

    heading(doc, 1, "6 Pipeline de telemetria")
    heading(doc, 2, "6.1 Etapas do processamento")
    numbered(doc, "A fonte ativa produz um TelemetrySample: mock, replay ou Assetto Corsa.")
    numbered(doc, "O TelemetryRuntime lê amostras em thread própria, normalmente a 60 Hz.")
    numbered(doc, "A amostra entra no buffer e no LapCollector.")
    numbered(doc, "O RuntimeState projeta a posição na pista e cria o frame de UI.")
    numbered(doc, "O event_bus emite processed_frame para WebSocket, gravação e modelo cognitivo.")
    numbered(doc, "O frontend recebe o frame, normaliza o payload e atualiza o store em intervalos controlados.")
    numbered(doc, "O mapa desenha carro, pista, histórico, oponentes e overlays sem depender de re-render React por frame.")
    heading(doc, 2, "6.2 Modelo canônico de amostra")
    para(
        doc,
        "TelemetrySample é a unidade central do sistema. Ela guarda posição de mundo, velocidade em km/h, "
        "yaw, progresso normalizado, volta, setor, throttle, brake, steering, marcha, RPM, acelerações, "
        "dados de pneus, suspensão, combustível, DRS, ABS, TC, temperatura e superfície quando disponíveis. "
        "Campos ausentes são aceitos; os módulos classificam a completude dos dados em vez de quebrar.",
    )

    heading(doc, 1, "7 Racing Line e análise de microsetores")
    heading(doc, 2, "7.1 Seleção da linha de referência")
    para(
        doc,
        "A Racing Line atual é baseada na melhor volta válida já completada pelo piloto. O sistema ignora "
        "a volta atual, rejeita voltas com poucas amostras, cobertura insuficiente do progresso, duração "
        "menor que 20 s ou maior que 900 s, e seleciona a menor duração válida como BEST_LAP.",
    )
    heading(doc, 2, "7.2 Linha visual")
    para(
        doc,
        "A linha visual não é desenhada a partir dos 50 microsetores. Ela usa visualLine.points, isto é, "
        "amostras reais da volta de referência ordenadas por splinePosition. Isso evita aparência de pista "
        "segmentada e evita que uma suavização artificial corte curvas como o S do Senna.",
    )
    heading(doc, 2, "7.3 Microsetores")
    para(
        doc,
        "Microsetores são divisões analíticas da volta. Por padrão, o usuário escolhe 20, 50 ou 100. "
        "Cada microsetor tem setor macro, início/fim de spline, velocidade média da referência, zona de "
        "freio, zona de aceleração, curvatura estimada, confiança e amostras usadas.",
    )
    heading(doc, 2, "7.4 Comparação com o piloto")
    para(
        doc,
        "A comparação calcula velocidade média do piloto, velocidade da referência, delta de velocidade, "
        "desvio médio de trajetória, freio/aceleração do piloto e da referência, tempo estimado e issue "
        "principal. As classes mais importantes são trajetória, freio cedo, freio tarde, aceleração tardia, "
        "baixa velocidade de contorno e baixa velocidade de saída.",
    )
    heading(doc, 2, "7.5 Overlay no mapa")
    caption(doc, "Quadro 7 - Modos da Racing Line no canvas")
    add_table(
        doc,
        ["Modo", "Função"],
        [
            ["LINE", "Mostra a linha fina, contínua e limpa da volta de referência."],
            ["DIAG", "Adiciona marcadores de problemas em cima da linha principal."],
            ["PEDAL", "Marca zonas de freio e aceleração da referência."],
            ["RAW", "Mostra pontos reais da volta de referência."],
            ["PTS", "Mostra pontos dos microsetores."],
            ["SEGS", "Mostra conexões segmentadas para debug."],
            ["SMTH", "Mostra linha suavizada apenas como camada de diagnóstico."],
            ["SPEED", "Mostra degradê por velocidade/freio: vermelho para frenagem/baixa velocidade e verde para alta velocidade."],
        ],
        widths=[2.5, 12.3],
    )
    source_note(doc, "Fonte: frontend/src/components/map/RacingLineOverlay.jsx.")

    heading(doc, 1, "8 Physics Fast Lap")
    heading(doc, 2, "8.1 Objetivo")
    para(
        doc,
        "Physics Fast Lap estima um perfil de velocidade alvo sobre a Racing Line existente. Ele não tenta "
        "criar uma linha ideal nova. A proposta é responder: dada esta linha real de referência, qual seria "
        "um perfil de velocidade fisicamente plausível para ela?",
    )
    heading(doc, 2, "8.2 Parâmetros atuais")
    caption(doc, "Quadro 8 - Parâmetros do modelo físico V1")
    add_table(
        doc,
        ["Parâmetro", "Valor atual", "Origem"],
        [
            ["Máximo lateral", "1,55 g", "CONFIG_DEFAULT"],
            ["Aceleração máxima", "4,2 m/s²", "CONFIG_DEFAULT"],
            ["Frenagem máxima", "10,0 m/s²", "CONFIG_DEFAULT"],
            ["Velocidade de reta", "95 m/s", "limite interno do modelo"],
            ["Massa", "Indisponível", "UNAVAILABLE"],
            ["Grip real da superfície", "Indisponível quando não vem do AC", "UNAVAILABLE"],
            ["Temperaturas", "Indisponíveis para cálculo de alvo", "UNAVAILABLE"],
        ],
        widths=[4.4, 4.8, 5.6],
    )
    source_note(doc, "Fonte: backend/core/physics_lap_model.py.")
    heading(doc, 2, "8.3 Como o alvo é calculado")
    numbered(doc, "Extrai amostras da visualLine da Racing Line.")
    numbered(doc, "Calcula distância acumulada entre pontos reais.")
    numbered(doc, "Estima curvatura com três pontos e suavização leve.")
    numbered(doc, "Calcula limite de velocidade por curvatura usando v = sqrt(a_lateral / curvatura).")
    numbered(doc, "Aplica passagem para frente respeitando aceleração máxima.")
    numbered(doc, "Aplica passagem para trás respeitando frenagem máxima.")
    numbered(doc, "Agrega resultado por microsetores e compara piloto, referência e alvo físico.")
    para(
        doc,
        "O relatório marca avisos importantes: physics_model_v1_simplified, uses_existing_reference_racing_line "
        "e not_a_full_assetto_corsa_physics_simulation. Portanto, o resultado deve ser lido como um alvo "
        "operacional inicial, não como simulação completa de dinâmica veicular.",
    )

    heading(doc, 1, "9 Inteligência de pilotagem e painéis de apoio")
    heading(doc, 2, "9.1 Race Coach")
    para(
        doc,
        "O Race Coach transforma segmentos da Racing Line em mensagens acionáveis. Ele agrupa microsetores "
        "consecutivos com o mesmo problema, calcula severidade e confiança, gera evidências e recomendações. "
        "Exemplos: atrasar levemente a frenagem, estabilizar antes do ápice, preparar a entrada para acelerar "
        "mais cedo ou comparar entrada/saída com a linha no mapa.",
    )
    heading(doc, 2, "9.2 Comparação ao vivo")
    para(
        doc,
        "O módulo de comparação ao vivo usa a volta atual, uma referência válida e histórico de oponentes. "
        "Ele compara velocidade, delta, trajetória, freio e aceleração por microsetor. Para oponentes, quando "
        "canais de pedal não existem, usa tendência de velocidade como inferência.",
    )
    heading(doc, 2, "9.3 Física do carro")
    para(
        doc,
        "O painel de física separa dados reais, inferidos e indisponíveis. Para o piloto, usa canais do AC "
        "quando disponíveis: velocidade, acelerações, pedais, pneus, suspensão, ambiente e estado do carro. "
        "Para oponentes, normalmente há velocidade e posição, com aceleração/frenagem inferidas por tendência.",
    )
    heading(doc, 2, "9.4 Modelo cognitivo")
    para(
        doc,
        "O CognitiveRuntime é leve e roda com throttle de emissão. Ele mede confiança, agressividade, suavidade "
        "e consistência a partir de variações de throttle, freio e direção. O DriverStateEstimator classifica "
        "estados como steady, overdriving, hesitating, confident_push ou unstable.",
    )

    heading(doc, 1, "10 Arquitetura do frontend")
    heading(doc, 2, "10.1 Layout")
    para(
        doc,
        "O Dashboard organiza a interface em mapa, painel direito e aba lateral esquerda de telemetria. "
        "A telemetria esquerda pode ser aberta ou ocultada. O painel direito possui abas Engineer, Debrief, "
        "Compare, Line e Physics. O mapa ocupa a região central e o controle de Replay fica no rodapé.",
    )
    heading(doc, 2, "10.2 Store e WebSocket")
    para(
        doc,
        "O useTelemetryStore mantém latestFrame, histórico, volta atual, volta anterior, voltas completadas, "
        "oponentes, eventos de coach, falas do engenheiro, estado cognitivo e modo de performance. O hook "
        "useTelemetryWS conecta em ws://hostname:8000/ws, normaliza mensagens e usa buffers pendentes para "
        "reduzir re-renders em React.",
    )
    heading(doc, 2, "10.3 Renderização do mapa")
    para(
        doc,
        "TrackRenderer usa canvas e requestAnimationFrame. O carro é suavizado com buffer temporal, atraso "
        "controlado de render, interpolação, extrapolação curta e reset em saltos grandes. O objetivo é reduzir "
        "travadas visuais sem falsificar a telemetria original.",
    )
    heading(doc, 2, "10.4 Modos de performance")
    caption(doc, "Quadro 9 - Modos de performance do frontend")
    add_table(
        doc,
        ["Modo", "Efeito"],
        [
            ["QUALITY", "Atualiza store com maior frequência e mantém histórico visual maior."],
            ["BALANCED", "Modo padrão; busca equilíbrio entre fluidez e custo de renderização."],
            ["PERFORMANCE", "Aumenta intervalos, reduz histórico e força Racing Line simples para aliviar o canvas."],
        ],
        widths=[3.4, 11.4],
    )
    source_note(doc, "Fonte: frontend/src/store/useTelemetryStore.ts e frontend/src/components/map/TrackRenderer.jsx.")

    heading(doc, 1, "11 Oponentes, gravação e persistência")
    heading(doc, 2, "11.1 Oponentes")
    para(
        doc,
        "O sistema possui um aplicativo Python para Assetto Corsa chamado ac_opponents_exporter. Ele envia "
        "snapshots UDP para 127.0.0.1:8765 com carros detectados, excluindo o player. O backend recebe, valida, "
        "remove dados obsoletos após alguns segundos e expõe /api/live/opponents.",
    )
    heading(doc, 2, "11.2 Gravação de sessão")
    para(
        doc,
        "RecordingRuntime assina eventos de telemetria do jogador, oponentes, splits, voltas e coaching. "
        "SessionRecorder grava JSONL em data/recordings, com metadata.json, player.jsonl, opponents.jsonl "
        "e events.jsonl. O gravador usa fila, batch, flush periódico e contabiliza frames descartados.",
    )
    heading(doc, 2, "11.3 Persistência local")
    para(
        doc,
        "O projeto contém bancos DuckDB para telemetria e feature store, caches FastF1, caches de geometria de "
        "pista e arquivos de debug. Os caches de pista evitam recalcular geometria KN5 a cada inicialização.",
    )

    heading(doc, 1, "12 Execução operacional")
    heading(doc, 2, "12.1 Inicialização segura")
    add_code_block(
        doc,
        [
            "cd C:\\Users\\Dokas\\Desktop\\f1-telemetry-analyzer\\backend",
            "$env:TELEMETRY_SOURCE='auto'",
            "$env:ASSETTO_AUTO_CONNECT='0'",
            "..\\.venv\\Scripts\\python.exe main.py",
            "",
            "cd C:\\Users\\Dokas\\Desktop\\f1-telemetry-analyzer\\frontend",
            "npm.cmd run dev",
        ],
    )
    source_note(doc, "Fonte: procedimento validado durante o desenvolvimento.")
    heading(doc, 2, "12.2 Conectar ao Assetto")
    numbered(doc, "Abra o Assetto Corsa e entre em uma sessão de pilotagem.")
    numbered(doc, "Confirme que o carro está carregado na pista, não apenas no menu.")
    numbered(doc, "Com o backend rodando, envie POST /api/live/source/assetto_corsa.")
    numbered(doc, "Verifique GET /api/live/source e GET /api/live/telemetry.")
    add_code_block(
        doc,
        [
            "Invoke-RestMethod -Method Post http://localhost:8000/api/live/source/assetto_corsa",
            "Invoke-RestMethod http://localhost:8000/api/live/source",
        ],
    )
    heading(doc, 2, "12.3 Desconectar antes de fechar ou reiniciar o Assetto")
    add_code_block(
        doc,
        [
            "Invoke-RestMethod -Method Post http://localhost:8000/api/live/source/mock",
        ],
    )
    para(
        doc,
        "Esse passo libera os handles de memória compartilhada do backend. Ele é recomendado antes de fechar "
        "ou reiniciar o Assetto, especialmente após o problema de crash em SharedMemoryWriter identificado "
        "durante a fase de integração.",
    )

    heading(doc, 1, "13 API do sistema")
    caption(doc, "Quadro 10 - Endpoints principais")
    add_table(
        doc,
        ["Endpoint", "Método", "Função"],
        [
            ["/health", "GET", "Status geral da aplicação."],
            ["/ws", "WebSocket", "Stream de telemetria e eventos para o frontend."],
            ["/api/live/source", "GET", "Mostra fonte ativa, reader, amostras e metadados."],
            ["/api/live/source/{source}", "POST", "Seleciona mock, replay ou assetto_corsa."],
            ["/api/live/telemetry", "GET", "Retorna frame atual e status de track/runtime."],
            ["/api/live/telemetry", "POST", "Ingere telemetria manual para testes."],
            ["/api/track/geometry", "GET", "Retorna geometria da pista para o mapa."],
            ["/api/track/cache", "GET/POST", "Consulta ou reconstrói cache de pista."],
            ["/api/live/comparison", "GET", "Comparação por microsetores."],
            ["/api/live/racing-line", "GET", "Payload da Racing Line e comparação."],
            ["/api/live/coach", "GET", "Relatório de coaching baseado na Racing Line."],
            ["/api/live/physics-fast-lap", "GET", "Perfil físico sobre a Racing Line existente."],
            ["/api/live/player-physics", "GET", "Física do piloto e oponentes."],
            ["/api/live/opponents", "GET", "Snapshot de oponentes recebidos via UDP."],
            ["/api/recording/status", "GET", "Status da gravação."],
            ["/api/recording/start", "POST", "Inicia gravação."],
            ["/api/recording/stop", "POST", "Para gravação."],
            ["/api/debug/performance", "GET", "Métricas de performance do backend."],
            ["/api/debug/projection", "GET", "Dados de debug da projeção espacial."],
            ["/api/debug/ac-shared-memory-full-inventory", "GET", "Inventário de campos AC."],
        ],
        widths=[6.2, 2.5, 7.0],
    )
    source_note(doc, "Fonte: backend/main.py.")

    heading(doc, 1, "14 Testes, validação e baseline")
    heading(doc, 2, "14.1 Testes de backend")
    para(
        doc,
        "Em 2 de junho de 2026, o backend foi validado com unittest discover no diretório backend/tests. "
        "Resultado: 72 testes executados em aproximadamente 1,24 s, todos aprovados.",
    )
    add_code_block(
        doc,
        [
            "Comando:",
            "..\\.venv\\Scripts\\python.exe -m unittest discover -s tests",
            "",
            "Resultado resumido:",
            "Ran 72 tests in 1.244s",
            "OK",
        ],
    )
    heading(doc, 2, "14.2 Build de frontend")
    para(
        doc,
        "O frontend foi validado com npm.cmd run build. O uso de npm.cmd evita bloqueio de política de execução "
        "do PowerShell sobre npm.ps1. Resultado: build Vite concluído com sucesso, 1.500 módulos transformados, "
        "bundle JavaScript de 351,72 kB e 105,86 kB gzip.",
    )
    add_code_block(
        doc,
        [
            "Comando:",
            "npm.cmd run build",
            "",
            "Resultado resumido:",
            "vite build",
            "1500 modules transformed",
            "built in 2.81s",
        ],
    )
    heading(doc, 2, "14.3 Cobertura funcional testada")
    bullet(doc, "AssettoAdapter não cria mmap nomeado quando a memória compartilhada ainda não existe.")
    bullet(doc, "TelemetrySourceManager não auto-conecta ao Assetto por padrão e exige amostra válida na conexão manual.")
    bullet(doc, "Racing Line seleciona a melhor volta válida, expõe visualLine real e rejeita voltas parciais.")
    bullet(doc, "Physics Fast Lap calcula distância, curvatura, passes de aceleração/frenagem e cache de perfil.")
    bullet(doc, "Race Coach, comparação, física, cognição, oponentes e gravação têm testes dedicados.")

    heading(doc, 1, "15 Troubleshooting")
    caption(doc, "Quadro 11 - Problemas comuns")
    add_table(
        doc,
        ["Sintoma", "Causa provável", "Ação recomendada"],
        [
            ["Assetto crasha em SharedMemoryWriter", "Aplicação externa abriu/criou memória compartilhada antes do simulador.", "Iniciar backend em modo seguro, abrir AC primeiro, conectar manualmente só em sessão e desconectar para mock antes de fechar AC."],
            ["Racing Line indisponível", "Nenhuma volta anterior completa e válida.", "Completar pelo menos uma volta limpa com cobertura suficiente."],
            ["Linha espelhada no mapa", "Transformação mundo -> canvas diferente da do carro.", "Garantir uso de X e Y=-Z, mesma função/convenção do mapPosition."],
            ["Linha parece pista segmentada", "Uso visual dos microsetores como linha principal.", "Usar visualLine.points como linha base e microsetores apenas como diagnósticos."],
            ["UI_RENDER_CRITICAL_FAILURE", "Carga alta de canvas, payload grande ou re-render excessivo.", "Usar modo PERFORMANCE, reduzir overlays, monitorar painel PERF e payloads."],
            ["Frontend não builda no PowerShell", "ExecutionPolicy bloqueia npm.ps1.", "Usar npm.cmd run build."],
            ["Sem oponentes", "Exporter do Assetto não instalado/aberto ou sessão sem carros válidos.", "Instalar app, habilitar no Content Manager e verificar /api/live/opponents."],
            ["Physics Fast Lap sem dados", "Racing Line ainda não está READY.", "Gerar uma volta válida de referência primeiro."],
        ],
        widths=[4.2, 5.2, 5.6],
    )
    source_note(doc)

    heading(doc, 1, "16 Limitações, decisões técnicas e roadmap")
    heading(doc, 2, "16.1 Limitações conhecidas")
    bullet(doc, "A Racing Line atual é a melhor volta válida, não uma linha ideal matemática.")
    bullet(doc, "Physics Fast Lap V1 é simplificado e não usa massa, aerodinâmica completa, pneus reais ou setup detalhado.")
    bullet(doc, "Oponentes normalmente não expõem pedais; freio/aceleração são inferidos por tendência de velocidade.")
    bullet(doc, "A geometria KN5 atual foi focada em Interlagos vhe_interlagos gp.")
    bullet(doc, "A interface ainda depende de polling para alguns relatórios analíticos, como Racing Line e Physics Fast Lap.")
    heading(doc, 2, "16.2 Decisões recentes importantes")
    bullet(doc, "Backend não auto-conecta ao Assetto por padrão: ASSETTO_AUTO_CONNECT=0.")
    bullet(doc, "Shared memory do AC só é aberta se todas as páginas já existem.")
    bullet(doc, "Linha visual da Racing Line usa amostras reais, sem suavização que corte curvas.")
    bullet(doc, "Gradiente de Racing Line é uma opção de visualização, não substitui diagnósticos por microsetor.")
    bullet(doc, "Left telemetry virou aba para abrir/ocultar e liberar espaço para leitura do mapa e painel direito.")
    heading(doc, 2, "16.3 Evolução recomendada")
    bullet(doc, "Criar botão de UI para conectar/desconectar Assetto sem precisar de chamada manual de API.")
    bullet(doc, "Evoluir Racing Line ideal composta a partir dos melhores microsetores de várias voltas.")
    bullet(doc, "Adicionar suavização com restrição de curvatura mínima, preservando amostras reais onde necessário.")
    bullet(doc, "Medir baseline antes/depois para qualquer otimização de canvas, payload ou re-render React.")
    bullet(doc, "Adicionar exportação de relatórios PDF/HTML por sessão.")
    bullet(doc, "Expandir suporte a pistas e validar o pipeline KN5 fora de Interlagos.")


def add_references(doc: Document):
    doc.add_page_break()
    heading(doc, 1, "Referências")
    refs = [
        "ASSOCIAÇÃO BRASILEIRA DE NORMAS TÉCNICAS. ABNT NBR 14724: informação e documentação - trabalhos acadêmicos - apresentação. Rio de Janeiro: ABNT, 2011.",
        "ASSOCIAÇÃO BRASILEIRA DE NORMAS TÉCNICAS. ABNT NBR 6023: informação e documentação - referências - elaboração. Rio de Janeiro: ABNT, 2018.",
        "ASSOCIAÇÃO BRASILEIRA DE NORMAS TÉCNICAS. ABNT NBR 6024: informação e documentação - numeração progressiva das seções de um documento - apresentação. Rio de Janeiro: ABNT, 2012.",
        "ASSOCIAÇÃO BRASILEIRA DE NORMAS TÉCNICAS. ABNT NBR 6027: informação e documentação - sumário - apresentação. Rio de Janeiro: ABNT, 2012.",
        "ASSOCIAÇÃO BRASILEIRA DE NORMAS TÉCNICAS. ABNT NBR 6028: informação e documentação - resumo, resenha e recensão - apresentação. Rio de Janeiro: ABNT, 2021.",
        "ASSOCIAÇÃO BRASILEIRA DE NORMAS TÉCNICAS. ABNT NBR 10520: informação e documentação - citações em documentos - apresentação. Rio de Janeiro: ABNT, 2023.",
        "F1 TELEMETRY ANALYZER. README.md, ARCHITECTURE.md, QUICKSTART.md, UDP_CAPTURE.md e CHANGELOG.md. Repositório local do projeto, 2026.",
        "F1 TELEMETRY ANALYZER. Código-fonte do backend e frontend. Repositório local do projeto, 2026.",
        "ASSOCIAÇÃO BRASILEIRA DE NORMAS TÉCNICAS. ABNT Catálogo. Disponível em: https://abntcatalogo.com.br/. Acesso em: 2 jun. 2026.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_paragraph_spacing(p, after=6, line=1.15, first_line=False)
        p.paragraph_format.left_indent = Cm(0)
        set_run(p.add_run(ref), size=11)


def add_appendices(doc: Document):
    doc.add_page_break()
    heading(doc, 1, "Apêndice A - Variáveis de ambiente")
    caption(doc, "Quadro 12 - Variáveis de ambiente operacionais")
    add_table(
        doc,
        ["Variável", "Valor padrão/uso", "Descrição"],
        [
            ["TELEMETRY_SOURCE", "auto", "Fonte inicial: auto, mock, replay ou assetto_corsa."],
            ["ASSETTO_AUTO_CONNECT", "0/false", "Quando falso, impede probe automático do Assetto."],
            ["ALLOW_REPLAY_FALLBACK", "false", "Permite replay quando auto não encontra fonte real."],
            ["TELEMETRY_DEBUG_REPLAY", "false", "Habilita replay fixture para debug."],
            ["TELEMETRY_RECORDING_ENABLED", "true", "Liga/desliga gravação local."],
            ["TELEMETRY_RECORDING_AUTO_START", "true", "Inicia gravação automaticamente no lifespan."],
            ["TELEMETRY_RECORDING_PLAYER_HZ", "20", "Frequência de gravação do player."],
            ["TELEMETRY_RECORDING_OPPONENTS_HZ", "20", "Frequência de gravação dos oponentes."],
            ["TELEMETRY_RECORDING_DIR", "data/recordings", "Diretório de saída das sessões."],
            ["VITE_API_URL", "http://localhost:8000", "Base URL usada pelo cliente Axios do frontend."],
        ],
        widths=[4.8, 4.0, 6.0],
    )
    source_note(doc)

    heading(doc, 1, "Apêndice B - Glossário técnico")
    caption(doc, "Quadro 13 - Termos essenciais")
    add_table(
        doc,
        ["Termo", "Explicação simples"],
        [
            ["Racing Line", "Linha que o piloto deve comparar visualmente. No estado atual, vem da melhor volta válida."],
            ["Microsetor", "Pequeno trecho da volta usado para análise estatística e diagnóstico."],
            ["Spline position", "Progresso normalizado da volta, de 0 a 1."],
            ["MapPosition", "Posição convertida para o canvas do mapa."],
            ["ProjectionEngine", "Módulo que encontra o ponto mais próximo do carro na pista."],
            ["KN5", "Arquivo de malha 3D usado pelo Assetto Corsa para pista/carros."],
            ["Shared memory", "Memória que o Assetto Corsa publica para outros programas lerem telemetria."],
            ["LapCollector", "Componente que identifica volta candidata e fechamento de volta."],
            ["Physics Fast Lap", "Perfil de velocidade alvo calculado sobre a Racing Line existente."],
            ["Race Coach", "Módulo que transforma perdas por microsetor em recomendações de pilotagem."],
        ],
        widths=[4.2, 10.8],
    )
    source_note(doc)


def build():
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_title_page(doc)
    add_summary(doc)
    add_static_toc(doc)
    add_body(doc)
    add_references(doc)
    add_appendices(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    out = build()
    print(out)
